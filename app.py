"""AI Resume Checker - Flask + scikit-learn backend.
Run:  pip install -r requirements.txt && python app.py  ->  http://127.0.0.1:5000
"""
import io, re
from flask import Flask, jsonify, render_template, request
from sklearn.feature_extraction.text import TfidfVectorizer, ENGLISH_STOP_WORDS
from sklearn.metrics.pairwise import cosine_similarity
from model import predict_roles
from insights import deep_insights

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 5 * 1024 * 1024

VERBS = set("led built developed designed implemented created managed improved increased reduced launched delivered optimized automated analyzed architected deployed mentored streamlined achieved drove owned shipped migrated trained engineered researched established generated saved".split())
WEAK = ["responsible for", "worked on", "helped with", "duties included", "assisted", "tasked with", "involved in"]
SECTIONS = {
    "summary": r"summary|objective|profile|about",
    "experience": r"experience|employment|work history|internship",
    "education": r"education|academic|degree",
    "skills": r"skills|technologies|tech stack|tools",
    "projects": r"projects?\b",
    "certifications": r"certif|awards|achievements",
}
SKILLS = set("python java javascript typescript c++ c# sql mysql postgresql mongodb react angular node flask django fastapi spring aws azure gcp docker kubernetes git linux tensorflow pytorch scikit-learn pandas numpy nlp opencv html css rest api machine learning deep learning data analysis excel tableau power bi iot arduino embedded".split())


def ocr_images(images):
    """OCR via Tesseract (needs the Tesseract binary installed, see README)."""
    try:
        import pytesseract
        return "\n".join(pytesseract.image_to_string(im) for im in images)
    except Exception as e:
        raise RuntimeError("OCR unavailable. Install Tesseract (see README). " + str(e))


def extract_text(file):
    name = (file.filename or "").lower()
    data = file.read()
    if name.endswith((".png", ".jpg", ".jpeg")):
        from PIL import Image
        return ocr_images([Image.open(io.BytesIO(data))])
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
        if len(text.strip()) < 80:  # likely a scanned PDF -> OCR fallback
            import pypdfium2 as pdfium
            pdf = pdfium.PdfDocument(data)
            text = ocr_images([pdf[i].render(scale=2).to_pil() for i in range(min(len(pdf), 3))])
        return text
    if name.endswith(".docx"):
        import docx
        return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
    if name.endswith(".odt"):
        import zipfile
        xml = zipfile.ZipFile(io.BytesIO(data)).read("content.xml").decode("utf-8", "ignore")
        return re.sub(r"<[^>]+>", " ", xml.replace("</text:p>", "\n"))
    text = data.decode("utf-8", errors="ignore")
    if name.endswith((".html", ".htm")):
        text = re.sub(r"<(script|style).*?</\1>", " ", text, flags=re.S | re.I)
        text = re.sub(r"<br\s*/?>|</(p|div|li|h\d)>", "\n", text, flags=re.I)
        text = re.sub(r"<[^>]+>", " ", text)
    if name.endswith(".rtf"):
        text = re.sub(r"\\par[d]?", "\n", text)
        text = re.sub(r"\\'[0-9a-f]{2}|\\[a-z]+-?\d* ?|[{}]", "", text)
    return text


def jd_match(resume, jd):
    """TF-IDF (1-2 grams) cosine similarity + coverage of top JD terms."""
    vec = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), sublinear_tf=True)
    m = vec.fit_transform([resume, jd])
    sim = float(cosine_similarity(m[0], m[1])[0][0])
    terms = vec.get_feature_names_out()
    weights = m[1].toarray()[0]
    top = [terms[i] for i in weights.argsort()[::-1][:25] if weights[i] > 0]
    rl = resume.lower()
    hit = [t for t in top if t in rl]
    miss = [t for t in top if t not in rl]
    score = round(min(100, sim * 140 * 0.4 + (len(hit) / max(1, len(top))) * 100 * 0.6))
    return score, hit, miss


def analyze(resume, jd=""):
    lines = [l for l in resume.splitlines() if l.strip()]
    words = len(resume.split())
    checks, fmt = [], 100
    if not re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", resume): fmt -= 15; checks.append(["bad", "No email address found."])
    if not re.search(r"\+?\d[\d\s().-]{8,}\d", resume): fmt -= 8; checks.append(["warn", "No phone number found."])
    if not re.search(r"linkedin\.com|github\.com", resume, re.I): fmt -= 6; checks.append(["warn", "Add a LinkedIn or GitHub link."])
    found = [k for k, p in SECTIONS.items() if any(len(l) < 40 and re.search(p, l, re.I) for l in lines)]
    for k in ("experience", "education", "skills"):
        if k not in found: fmt -= 10; checks.append(["bad", f'Missing clear "{k}" section heading.'])
    if words < 250: fmt -= 12; checks.append(["warn", f"Only {words} words - too short (aim 400-700)."])
    elif words > 900: fmt -= 10; checks.append(["warn", f"{words} words - too long, trim to 1-2 pages."])
    bullets = [l for l in lines if re.match(r"\s*[-•*▪·]", l)] or [l for l in lines if len(l) > 45]
    first = lambda l: (re.findall(r"[a-z]+", l.lower()) or [""])[0]
    verb_b = [l for l in bullets if first(l) in VERBS]
    quant = [l for l in bullets if re.search(r"\d+(\.\d+)?\s?(%|x|k|m|\+)|\$\s?\d|\b\d{2,}\b", l, re.I)]
    weak = [w for w in WEAK if w in resume.lower()]
    n = max(1, len(bullets))
    impact = max(5, min(100, round(len(verb_b) / n * 45 + len(quant) / n * 55 + 10 - (8 if weak else 0))))
    impact_notes = [f"{len(verb_b)}/{len(bullets)} bullets start with a strong action verb.",
                    f"{len(quant)}/{len(bullets)} bullets contain metrics (target 50%+)."]
    if weak: impact_notes.append("Weak phrases: " + ", ".join(weak))
    skills = sorted(s for s in SKILLS if re.search(r"(?<![\w+#])" + re.escape(s) + r"(?![\w+#])", resume.lower()))
    match, hit, miss = (None, [], [])
    if len(jd) > 40:
        match, hit, miss = jd_match(resume, jd)
        if miss: checks.append(["warn", f"Missing {len(miss)} high-weight job keywords."])
    fmt = max(0, fmt)
    overall = round(fmt * .25 + impact * .3 + match * .45) if match is not None else round(fmt * .5 + impact * .5)
    if not checks: checks.append(["ok", "Structure looks ATS-friendly."])
    return dict(overall=overall, format=fmt, impact=impact, match=match, skills=skills, hit=hit,
                miss=miss, roles=predict_roles(resume), checks=checks, impact_notes=impact_notes, sections=found, words=words, **deep_insights(resume, lines, bullets, miss, VERBS))


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/analyze", methods=["POST"])
def analyze_route():
    text = request.form.get("resume", "")
    f = request.files.get("file")
    if f and f.filename:
        try:
            text = extract_text(f)
        except Exception as e:
            return jsonify(error=str(e)), 400
    if len(text.strip()) < 80:
        return jsonify(error="Resume text too short or unreadable."), 400
    return jsonify(analyze(text, request.form.get("jd", "")))


if __name__ == "__main__":
    app.run(debug=True)
